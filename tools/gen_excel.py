# -*- coding: utf-8 -*-
"""Generate the interactive Excel dashboard (Arabic, blue/white theme)."""
import sys
import datetime
import calendar
from collections import Counter, defaultdict
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else ("/tmp/xl_src.xls" if __import__("os").path.exists("/tmp/xl_src.xls") else "/tmp/xl_src.xlsx")
_proot = __import__("pathlib").Path(__file__).resolve().parent.parent
OUT = sys.argv[2] if len(sys.argv) > 2 else str(_proot / "EPC-Interactive-Dashboard.xlsx")
_llogo = _proot / "assets" / "logo.png"
LOGO = str(_llogo) if _llogo.exists() else "/tmp/logo.png"
BUF = 2000  # buffered rows in Data sheet for future pastes

BLUE = "1976D2"; BLUE_D = "0D47A1"; LIGHT = "E3F2FD"; LINE = "DBE4EF"
WHITE = "FFFFFF"; MUT = "5F7288"; YELLOW = "FFF9C4"

thin = Side(style="thin", color=LINE)
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------- extract from ERP export ----------
if SRC.lower().endswith(".xls"):
    import xlrd
    try:
        bk = xlrd.open_workbook(SRC, encoding_override="cp1256")  # ERP exports: Arabic Win-1256
    except Exception:
        bk = xlrd.open_workbook(SRC)
    sh = bk.sheet_by_name("0040_qa") if "0040_qa" in bk.sheet_names() else bk.sheet_by_index(0)
    rows = []
    for i in range(sh.nrows):
        vals = []
        for j in range(sh.ncols):
            v = sh.cell_value(i, j)
            if sh.cell_type(i, j) == xlrd.XL_CELL_DATE:
                v = xlrd.xldate.xldate_as_datetime(v, bk.datemode)
            vals.append(v)
        rows.append(vals)
else:
    wb_src = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    ws_src = wb_src["0040_qa"] if "0040_qa" in wb_src.sheetnames else wb_src.worksheets[0]
    rows = [list(r) for r in ws_src.iter_rows(values_only=True)]
    wb_src.close()
hdr = {str(v).strip().lower(): i for i, v in enumerate(rows[0]) if v is not None}
def g(r, n):
    i = hdr.get(n)
    return r[i] if i is not None and i < len(r) else None
def f(x):
    try: return float(x)
    except Exception: return 0.0
def s(x):
    if x is None: return ""
    if isinstance(x, float) and x.is_integer(): return str(int(x))
    return str(x).strip()

data = []
for r in rows[1:]:
    d = g(r, "prdate")
    if not hasattr(d, "day") or not s(g(r, "packno")):
        continue
    q = f(g(r, "packqty"))
    ct = q * f(g(r, "cond_kg")) / 1000.0
    st = q * f(g(r, "screen_kgm")) / 1000.0
    at = q * f(g(r, "armour_kgm")) / 1000.0
    data.append(dict(d=d.day, pk=s(g(r, "packno")), wo=s(g(r, "wo")), cu=s(g(r, "customer")) or "?",
                     fam=s(g(r, "prodfamily")) or "?", vs=(s(g(r, "volt")) + " " + s(g(r, "size"))).strip(),
                     mac=s(g(r, "maccode")) or "?", pt=s(g(r, "packtype")) or "?", qc=s(g(r, "qcstatus")) or "P",
                     con=s(g(r, "conductor")) or "?", exp="Y" if s(g(r, "export")).upper() == "Y" else "N",
                     q=round(q), ct=ct, st=st, at=at, tt=ct + st + at))
ym = Counter((2026, 10))  # replaced below
ym = Counter([(2026, 10)])
dates = Counter()
for r in rows[1:]:
    d = g(r, "prdate")
    if hasattr(d, "year"):
        dates[(d.year, d.month)] += 1
(year, month) = dates.most_common(1)[0][0]
MDAYS = calendar.monthrange(year, month)[1]
AR = {1:"يناير",2:"فبراير",3:"مارس",4:"أبريل",5:"مايو",6:"يونيو",7:"يوليو",8:"أغسطس",9:"سبتمبر",10:"أكتوبر",11:"نوفمبر",12:"ديسمبر"}
fams = ["MV", "LV", "HV", "OH", "BC", "CC", "IC"]
custs = [c for c, _ in Counter(x["cu"] for x in data).most_common()]
macs = [m for m, _ in Counter(x["mac"] for x in data).most_common()]

wb = Workbook()

# ================= الداشبورد =================
ws = wb.active
ws.title = "الداشبورد"
ws.sheet_view.rightToLeft = True
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVWX", [16] * 24):
    ws.column_dimensions[col].width = w

img = XLImage(LOGO)
img.height = 60; img.width = 107
ws.add_image(img, "A1")
ws["C2"] = "لوحة متابعة الإنتاج التفاعلية — Energya Power Cables"
ws["C2"].font = Font(size=16, bold=True, color=BLUE_D)
ws["C3"] = f"{AR[month]} {year} · البيانات حتى يوم {max(x['d'] for x in data)} · عدّل القوائم المنسدلة وكل الأرقام والرسوم تتحدث فوراً"
ws["C3"].font = Font(size=10, color=MUT)

# --- selectors ---
sels = [("اليوم", "B2"), ("العائلة", "B3"), ("السوق", "B4"), ("الموصل", "B5"),
        ("الجودة", "B6"), ("العميل", "B7"), ("الماكينة", "B8")]
# put label in col A row 5..11 and dropdown in col B
for i, (label, _) in enumerate(sels):
    r = 5 + i
    ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=10, color=BLUE_D)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    ws.cell(row=r, column=1).border = BORD
    c = ws.cell(row=r, column=2, value="الكل")
    c.fill = PatternFill("solid", fgColor=YELLOW)
    c.border = BORD
    c.font = Font(bold=True, size=10)
    c.alignment = Alignment(horizontal="center")
SEL = {name: f"$B${5+i}" for i, (name, _) in enumerate(sels)}  # اليوم..الماكينة

dv_day = DataValidation(type="list", formula1='"الكل,1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,22,23,24,25,26,27,28,29,30,31"', allow_blank=True)
dv_fam = DataValidation(type="list", formula1="=قوائم!$A$2:$A$9", allow_blank=True)
dv_mkt = DataValidation(type="list", formula1="=قوائم!$E$2:$E$4", allow_blank=True)
dv_con = DataValidation(type="list", formula1="=قوائم!$F$2:$F$4", allow_blank=True)
dv_qc  = DataValidation(type="list", formula1="=قوائم!$D$2:$D$5", allow_blank=True)
dv_cus = DataValidation(type="list", formula1=f"=قوائم!$B$2:$B${len(custs)+1}", allow_blank=True)
dv_mac = DataValidation(type="list", formula1=f"=قوائم!$C$2:$C${len(macs)+1}", allow_blank=True)
for dv, cell in [(dv_day, "B5"), (dv_fam, "B6"), (dv_mkt, "B7"), (dv_con, "B8"), (dv_qc, "B9"), (dv_cus, "B10"), (dv_mac, "B11")]:
    ws.add_data_validation(dv)
    dv.add(ws[cell])
SEL = {"day": "$B$5", "fam": "$B$6", "mkt": "$B$7", "con": "$B$8", "qc": "$B$9", "cust": "$B$10", "mac": "$B$11"}

# --- KPI cards ---
D = "'البيانات'!"
X = f"{D}$X$2:$X${BUF+1}"
P = f"{D}$P$2:$P${BUF+1}"
L = f"{D}$L$2:$L${BUF+1}"
kpis = [
    ("الإنتاج الفعلي (طن)", f"=SUMPRODUCT({X},{P})", "#,##0.0"),
    ("الخطة (طن)", "=IF('الخطة'!$B$1>0,'الخطة'!$B$1,0)", "#,##0"),
    ("نسبة التحقيق", "=IF($E$15>0,$C$15/$E$15,0)", "0.0%"),
    ("الفاضل (طن)", "=IF($E$15>0,MAX(0,$E$15-$C$15),0)", "#,##0.0"),
    ("أيام متبقية", "=MAX(0,'الخطة'!$B$4-MAX('البيانات'!$A$2:$A$2000))", "0"),
    ("مطلوب طن/يوم", "=IF($K$15>0,$I$15/$K$15,0)", "#,##0.0"),
    ("المعدل اليومي الفعلي", "=IF($G$18>0,$C$15/$G$18,0)", "#,##0.0"),
    ("توقع نهاية الشهر", "=$C$18*'الخطة'!$B$4", "#,##0"),
    ("أيام إنتاج", f"=COUNTIF('حسابات'!$B$2:$B${1+MDAYS},\">0\")", "0"),
    ("إجمالي كم", f"=SUMPRODUCT({X},{L})/1000", "#,##0.0"),
    ("طرود", f"=SUMPRODUCT({X})", "#,##0"),
    ("حصة التصدير", f"=IF($C$15>0,SUMPRODUCT({X},--({D}$K$2:$K${BUF+1}=\"Y\"),{P})/$C$15,0)", "0.0%"),
]
ws["C13"] = "مؤشرات الشهر"; ws["C13"].font = Font(bold=True, size=12, color=BLUE_D)
for i, (label, formula, fmt) in enumerate(kpis):
    col = 3 + (i % 6) * 2
    row = 14 + (i // 6) * 3
    lc = ws.cell(row=row, column=col, value=label)
    lc.font = Font(size=9, bold=True, color=BLUE_D)
    lc.fill = PatternFill("solid", fgColor=LIGHT)
    lc.border = BORD
    lc.alignment = Alignment(horizontal="center")
    vc = ws.cell(row=row + 1, column=col, value=formula)
    vc.number_format = fmt
    vc.font = Font(size=15, bold=True, color=BLUE_D)
    vc.fill = PatternFill("solid", fgColor=WHITE)
    vc.border = BORD
    vc.alignment = Alignment(horizontal="center")
# data bar on achievement %
ws.conditional_formatting.add("G15", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=BLUE, showValue=True))
ws["C20"] = "💡 غيّر أي قائمة منسدلة فوق — كل المؤشرات والرسوم بتتحدث لحظياً. الخطة تتعدل من ورقة «الخطة»."
ws["C20"].font = Font(size=10, italic=True, color=MUT)

# ================= البيانات =================
wd = wb.create_sheet("البيانات")
wd.sheet_view.rightToLeft = True
cols = [("اليوم", "d"), ("الطرد", "pk"), ("أمر التشغيل", "wo"), ("العميل", "cu"), ("العائلة", "fam"),
        ("المواصفة", "vs"), ("الماكينة", "mac"), ("نوع التعبئة", "pt"), ("حالة QC", "qc"),
        ("الموصل", "con"), ("السوق", "exp"), ("أمتار", "q"), ("طن موصلات", "ct"), ("طن Screen", "st"),
        ("طن Armour", "at"), ("الطن الكلي", "tt")]
for j, (name, _) in enumerate(cols, start=1):
    c = wd.cell(row=1, column=j, value=name)
    c.font = Font(bold=True, color=WHITE, size=10)
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.border = BORD
    c.alignment = Alignment(horizontal="center")
for i, rec in enumerate(data, start=2):
    for j, (_, key) in enumerate(cols, start=1):
        v = rec[key]
        c = wd.cell(row=i, column=j, value=v)
        c.border = BORD
        if isinstance(v, float):
            c.number_format = "#,##0.000"
# helper match columns Q..X
help_defs = [
    ("مطابقة اليوم", f'=--($A2=IF(الداشبورد!{SEL["day"]}="الكل",$A2,الداشبورد!{SEL["day"]}))'),
    ("مطابقة العائلة", f'=--($E2=IF(الداشبورد!{SEL["fam"]}="الكل",$E2,الداشبورد!{SEL["fam"]}))'),
    ("مطابقة السوق", f'=--($K2=IF(الداشبورد!{SEL["mkt"]}="الكل",$K2,IF(الداشبورد!{SEL["mkt"]}="تصدير","Y","N")))'),
    ("مطابقة الموصل", f'=--($J2=IF(الداشبورد!{SEL["con"]}="الكل",$J2,IF(الداشبورد!{SEL["con"]}="AL","A","C")))'),
    ("مطابقة الجودة", f'=--($I2=IF(الداشبورد!{SEL["qc"]}="الكل",$I2,IF(الداشبورد!{SEL["qc"]}="بدون حالة","P",الداشبورد!{SEL["qc"]})))'),
    ("مطابقة العميل", f'=--($D2=IF(الداشبورد!{SEL["cust"]}="الكل",$D2,الداشبورد!{SEL["cust"]}))'),
    ("مطابقة الماكينة", f'=--($G2=IF(الداشبورد!{SEL["mac"]}="الكل",$G2,الداشبورد!{SEL["mac"]}))'),
]
# NOTE: unquoted arabic sheet refs fixed below via replace
for j, (name, formula) in enumerate(help_defs, start=17):  # Q=17
    c = wd.cell(row=1, column=j, value=name)
    c.font = Font(bold=True, color=WHITE, size=9)
    c.fill = PatternFill("solid", fgColor="64748B")
    c.border = BORD
    formula = formula.replace("الداشبورد!", "'الداشبورد'!")
    for i in range(2, BUF + 2):
        wd.cell(row=i, column=j, value=formula.replace("$A2", f"$A{i}").replace("$E2", f"$E{i}").replace("$K2", f"$K{i}").replace("$J2", f"$J{i}").replace("$I2", f"$I{i}").replace("$D2", f"$D{i}").replace("$G2", f"$G{i}"))
c = wd.cell(row=1, column=24, value="مطابق لكل الفلاتر")
c.font = Font(bold=True, color=WHITE, size=9)
c.fill = PatternFill("solid", fgColor=BLUE_D)
c.border = BORD
for i in range(2, BUF + 2):
    wd.cell(row=i, column=24, value=f"=Q{i}*R{i}*S{i}*T{i}*U{i}*V{i}*W{i}*--($C{i}<>\"\")")
wd.freeze_panes = "A2"
wd.auto_filter.ref = f"A1:P{len(data)+1}"
for col, w in zip(["A","B","C","D","E","F","G","H","I","J","K","L","M","N","O","P"], [7,11,12,30,9,18,10,14,9,8,7,9,11,11,11,11]):
    wd.column_dimensions[col].width = w

# ================= حسابات =================
wc = wb.create_sheet("حسابات")
wc.sheet_view.rightToLeft = True
Xr = f"'البيانات'!$X$2:$X${BUF+1}"
Pr = f"'البيانات'!$P$2:$P${BUF+1}"
Er = f"'البيانات'!$E$2:$E${BUF+1}"
Jr = f"'البيانات'!$J$2:$J${BUF+1}"
Mr = f"'البيانات'!$M$2:$M${BUF+1}"
Nr = f"'البيانات'!$N$2:$N${BUF+1}"
Or_ = f"'البيانات'!$O$2:$O${BUF+1}"
Kr = f"'البيانات'!$K$2:$K${BUF+1}"
Ir = f"'البيانات'!$I$2:$I${BUF+1}"
Dr = f"'البيانات'!$D$2:$D${BUF+1}"
Gr = f"'البيانات'!$G$2:$G${BUF+1}"
Ar_ = f"'البيانات'!$A$2:$A${BUF+1}"
POWER = f"((({Er}=\"LV\")+({Er}=\"MV\")+({Er}=\"HV\")+({Er}=\"OH\"))>0)"
# daily table A:day B:total C:power D:telecom E:cum F:plan cum
wc["A1"], wc["B1"], wc["C1"], wc["D1"], wc["E1"], wc["F1"] = "اليوم", "إجمالي اليوم", "Power", "Telecom", "تراكمي فعلي", "تراكمي الخطة"
for j in range(1, 7):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
for d in range(1, MDAYS + 1):
    r = 1 + d
    wc.cell(row=r, column=1, value=d)
    wc.cell(row=r, column=2, value=f"=SUMPRODUCT({Xr},--({Ar_}=$A{r}),{Pr})")
    wc.cell(row=r, column=3, value=f"=SUMPRODUCT({Xr},--({Ar_}=$A{r}),--{POWER},{Pr})")
    wc.cell(row=r, column=4, value=f"=$B{r}-$C{r}")
    wc.cell(row=r, column=5, value=f"=SUM($B$2:$B{r})")
    wc.cell(row=r, column=6, value=f"=IF('الخطة'!$B$1>0,'الخطة'!$B$1/'الخطة'!$B$4*$A{r},0)")
# family table H:I
wc["H1"], wc["I1"] = "العائلة", "طن"
for j in (8, 9):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
for i, fam in enumerate(fams):
    r = 2 + i
    wc.cell(row=r, column=8, value=fam)
    wc.cell(row=r, column=9, value=f"=SUMPRODUCT({Xr},--({Er}=$H{r}),{Pr})")
# materials L:M
wc["L1"], wc["M1"] = "المادة", "طن"
for j in (12, 13):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
mats = [("AL ألومنيوم", f"=SUMPRODUCT({Xr},--({Jr}=\"A\"),{Mr})"),
        ("CU نحاس", f"=SUMPRODUCT({Xr},--({Jr}=\"C\"),{Mr})"),
        ("Screen", f"=SUMPRODUCT({Xr},{Nr})"),
        ("Armour", f"=SUMPRODUCT({Xr},{Or_})")]
for i, (name, formula) in enumerate(mats):
    wc.cell(row=2 + i, column=12, value=name)
    wc.cell(row=2 + i, column=13, value=formula)
# market P:Q
wc["P1"], wc["Q1"] = "السوق", "طن"
for j in (16, 17):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
wc["P2"], wc["Q2"] = "تصدير", f"=SUMPRODUCT({Xr},--({Kr}=\"Y\"),{Pr})"
wc["P3"], wc["Q3"] = "محلي", f"=SUMPRODUCT({Xr},--({Kr}=\"N\"),{Pr})"
# qc T:U
wc["T1"], wc["U1"] = "حالة QC", "طرود"
for j in (20, 21):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
qcs = [("OK", f"=SUMPRODUCT({Xr},--({Ir}=\"OK\"))"), ("Hold", f"=SUMPRODUCT({Xr},--({Ir}=\"Hold\"))"), ("بدون حالة", f"=SUMPRODUCT({Xr},--({Ir}=\"P\"))")]
for i, (name, formula) in enumerate(qcs):
    wc.cell(row=2 + i, column=20, value=name)
    wc.cell(row=2 + i, column=21, value=formula)
# customers top X:Y
wc["X1"], wc["Y1"] = "العميل", "طن"
for j in (24, 25):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
for i, cu in enumerate(custs[:12]):
    wc.cell(row=2 + i, column=24, value=cu)
    wc.cell(row=2 + i, column=25, value=f"=SUMPRODUCT({Xr},--({Dr}=$X{2+i}),{Pr})")
# machines AB:AC
wc["AB1"], wc["AC1"] = "الماكينة", "طن"
for j in (28, 29):
    wc.cell(row=1, column=j).font = Font(bold=True, color=WHITE, size=9)
    wc.cell(row=1, column=j).fill = PatternFill("solid", fgColor=BLUE)
for i, m in enumerate(macs[:10]):
    wc.cell(row=2 + i, column=28, value=m)
    wc.cell(row=2 + i, column=29, value=f"=SUMPRODUCT({Xr},--({Gr}=$AB{2+i}),{Pr})")

# ================= الخطة =================
wp = wb.create_sheet("الخطة")
wp.sheet_view.rightToLeft = True
wp["A1"], wp["B1"] = "إجمالي خطة الشهر (طن)", 4442
wp["A2"], wp["B2"] = "خطة الألومنيوم AL (طن)", 3139
wp["A3"], wp["B3"] = "خطة النحاس CU (طن)", 1303
wp["A4"], wp["B4"] = "عدد أيام الشهر", MDAYS
wp["A5"], wp["B5"] = "الشهر", f"{AR[month]} {year}"
for r in range(1, 6):
    wp.cell(row=r, column=1).font = Font(bold=True, color=BLUE_D, size=10)
    wp.cell(row=r, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    wp.cell(row=r, column=1).border = BORD
    wp.cell(row=r, column=2).border = BORD
    wp.cell(row=r, column=2).font = Font(bold=True, size=11)
    if r <= 4:
        wp.cell(row=r, column=2).fill = PatternFill("solid", fgColor=YELLOW)
wp["A7"] = "✏️ عدّل الأرقام الصفراء بس — الداشبورد كله يتحدث لوحده."
wp["A7"].font = Font(italic=True, color=MUT, size=10)
wp.column_dimensions["A"].width = 30
wp.column_dimensions["B"].width = 16

# ================= قوائم =================
wl = wb.create_sheet("قوائم")
lists = {"A": ["الكل"] + fams, "B": ["الكل"] + custs, "C": ["الكل"] + macs,
         "D": ["الكل", "OK", "Hold", "بدون حالة"], "E": ["الكل", "تصدير", "محلي"], "F": ["الكل", "AL", "CU"]}
for col, vals in lists.items():
    for i, v in enumerate(vals, start=1):
        wl[f"{col}{i}"] = v
wl.sheet_state = "hidden"
wc.sheet_state = "hidden"

# ================= charts on dashboard =================
def style(ch, title):
    ch.title = title
    ch.height = 9
    ch.width = 16
    ch.style = 10
    return ch

ch1 = LineChart()
cats = Reference(wc, min_col=1, min_row=2, max_row=1 + MDAYS)
s_act = Reference(wc, min_col=5, min_row=1, max_row=1 + MDAYS)
s_plan = Reference(wc, min_col=6, min_row=1, max_row=1 + MDAYS)
ch1.add_data(s_act, titles_from_data=True)
ch1.add_data(s_plan, titles_from_data=True)
ch1.set_categories(cats)
ch1.series[0].graphicalProperties.line.solidFill = BLUE
ch1.series[0].graphicalProperties.line.width = 28000
ch1.series[1].graphicalProperties.line.solidFill = "D97706"
ch1.series[1].graphicalProperties.line.dashStyle = "dash"
style(ch1, "الإنتاج التراكمي مقابل الخطة")
ws.add_chart(ch1, "C22")

ch2 = BarChart()
ch2.type = "col"; ch2.grouping = "stacked"; ch2.overlap = 100
pow_ref = Reference(wc, min_col=3, min_row=1, max_row=1 + MDAYS)
tel_ref = Reference(wc, min_col=4, min_row=1, max_row=1 + MDAYS)
ch2.add_data(pow_ref, titles_from_data=True)
ch2.add_data(tel_ref, titles_from_data=True)
ch2.set_categories(cats)
ch2.series[0].graphicalProperties.solidFill = BLUE
ch2.series[1].graphicalProperties.solidFill = "F59E0B"
style(ch2, "الإنتاج اليومي Power / Telecom")
ws.add_chart(ch2, "O22")

def colorize(series, colors):
    for i, c in enumerate(colors):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = c
        series.data_points.append(dp)

ch3 = DoughnutChart()
fam_vals = Reference(wc, min_col=9, min_row=1, max_row=1 + len(fams))
fam_cats = Reference(wc, min_col=8, min_row=2, max_row=1 + len(fams))
ch3.add_data(fam_vals, titles_from_data=True)
ch3.set_categories(fam_cats)
colorize(ch3.series[0], ["1976D2", "0F9D58", "7C3AED", "DC2626", "EA580C", "CA8A04", "64748B"])
style(ch3, "العائلات (طن)")
ws.add_chart(ch3, "C40")

ch4 = PieChart()
mk_vals = Reference(wc, min_col=17, min_row=1, max_row=3)
mk_cats = Reference(wc, min_col=16, min_row=2, max_row=3)
ch4.add_data(mk_vals, titles_from_data=True)
ch4.set_categories(mk_cats)
colorize(ch4.series[0], ["0F9D58", "1976D2"])
style(ch4, "تصدير / محلي")
ws.add_chart(ch4, "O40")

ch5 = BarChart()
ch5.type = "bar"
mt_vals = Reference(wc, min_col=13, min_row=1, max_row=5)
mt_cats = Reference(wc, min_col=12, min_row=2, max_row=5)
ch5.add_data(mt_vals, titles_from_data=True)
ch5.set_categories(mt_cats)
ch5.series[0].graphicalProperties.solidFill = "8FA3BD"
style(ch5, "المواد (طن)")
ws.add_chart(ch5, "C58")

ch6 = DoughnutChart()
qc_vals = Reference(wc, min_col=21, min_row=1, max_row=4)
qc_cats = Reference(wc, min_col=20, min_row=2, max_row=4)
ch6.add_data(qc_vals, titles_from_data=True)
ch6.set_categories(qc_cats)
colorize(ch6.series[0], ["0B8A4B", "C77700", "64748B"])
style(ch6, "حالة الجودة (طرود)")
ws.add_chart(ch6, "O58")

ch7 = BarChart()
ch7.type = "bar"
cu_vals = Reference(wc, min_col=25, min_row=1, max_row=13)
cu_cats = Reference(wc, min_col=24, min_row=2, max_row=13)
ch7.add_data(cu_vals, titles_from_data=True)
ch7.set_categories(cu_cats)
ch7.series[0].graphicalProperties.solidFill = BLUE
style(ch7, "أعلى العملاء (طن)")
ch7.height = 10
ws.add_chart(ch7, "C76")

ch8 = BarChart()
ch8.type = "bar"
ma_vals = Reference(wc, min_col=29, min_row=1, max_row=11)
ma_cats = Reference(wc, min_col=28, min_row=2, max_row=11)
ch8.add_data(ma_vals, titles_from_data=True)
ch8.set_categories(ma_cats)
ch8.series[0].graphicalProperties.solidFill = "0D9488"
style(ch8, "الماكينات (طن)")
ch8.height = 10
ws.add_chart(ch8, "O76")

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "| packs:", len(data), "| month:", year, month, "| customers:", len(custs), "| machines:", len(macs))
